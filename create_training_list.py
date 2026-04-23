"""
N梱包 教育完了リスト 生成スクリプト

メンバーと教育項目をもとに Excel シートを生成します。
セルに ○ を入力すると緑、× を入力すると赤に自動着色されます。

使用方法:
    python create_training_list.py
    → N梱包_教育完了リスト.xlsx が生成されます
"""

from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side,
    GradientFill
)
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule
from openpyxl.utils import get_column_letter

OUTPUT_FILE = "N梱包_教育完了リスト.xlsx"

MEMBERS = ["館澤", "若園", "小栗", "森", "近藤", "佐々木"]

ITEMS = [
    "当日即出荷",
    "当日残１即出荷",
    "翌日出荷",
    "複数口梱包",
    "冷蔵品梱包",
    "冷凍品梱包",
    "割れ物梱包",
    "大型品梱包",
    "ラベル貼付",
    "伝票作成",
    "検品・数量確認",
    "仕分け作業",
    "入荷処理",
    "在庫管理入力",
    "返品処理",
    "お届け日指定便",
    "時間指定便",
    "代引き処理",
    "クレーム対応フロー",
    "緊急出荷対応",
]

# 色定義
COLOR_HEADER_BG    = "1F4E79"   # 濃紺（ヘッダー背景）
COLOR_HEADER_FONT  = "FFFFFF"   # 白（ヘッダー文字）
COLOR_TITLE_BG     = "2E75B6"   # 青（タイトル行背景）
COLOR_ROW_ODD      = "DEEAF1"   # 薄青（奇数行）
COLOR_ROW_EVEN     = "FFFFFF"   # 白（偶数行）
COLOR_COMPLETE     = "C6EFCE"   # 薄緑（○）
COLOR_COMPLETE_FT  = "276221"   # 濃緑（○ 文字）
COLOR_INCOMPLETE   = "FFC7CE"   # 薄赤（×）
COLOR_INCOMPLETE_FT= "9C0006"   # 濃赤（× 文字）
COLOR_BORDER       = "9DC3E6"   # 薄青（罫線）


def thin_border(color=COLOR_BORDER):
    side = Side(style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def header_border():
    side = Side(style="medium", color="FFFFFF")
    return Border(left=side, right=side, top=side, bottom=side)


def apply_conditional_formatting(ws, data_range):
    """○ → 緑、× → 赤 の条件付き書式を設定"""
    # ○ のとき緑
    green_fill = PatternFill(start_color=COLOR_COMPLETE, end_color=COLOR_COMPLETE, fill_type="solid")
    green_font = Font(color=COLOR_COMPLETE_FT, bold=True)
    ds_green = DifferentialStyle(fill=green_fill, font=green_font)
    rule_green = Rule(type="containsText", operator="containsText", text="○", dxf=ds_green)
    rule_green.formula = [f'NOT(ISERROR(SEARCH("○",{data_range.split(":")[0]})))']
    ws.conditional_formatting.add(data_range, rule_green)

    # × のとき赤
    red_fill = PatternFill(start_color=COLOR_INCOMPLETE, end_color=COLOR_INCOMPLETE, fill_type="solid")
    red_font = Font(color=COLOR_INCOMPLETE_FT, bold=True)
    ds_red = DifferentialStyle(fill=red_fill, font=red_font)
    rule_red = Rule(type="containsText", operator="containsText", text="×", dxf=ds_red)
    rule_red.formula = [f'NOT(ISERROR(SEARCH("×",{data_range.split(":")[0]})))']
    ws.conditional_formatting.add(data_range, rule_red)


def create_sheet(wb):
    ws = wb.active
    ws.title = "教育完了リスト"

    # ── 列幅の設定 ──────────────────────────────
    ws.column_dimensions["A"].width = 5   # No.
    ws.column_dimensions["B"].width = 22  # 項目名
    for col_idx in range(len(MEMBERS)):
        col_letter = get_column_letter(col_idx + 3)
        ws.column_dimensions[col_letter].width = 10

    # ── タイトル行 (row 1) ────────────────────────
    last_col = get_column_letter(2 + len(MEMBERS))
    ws.merge_cells(f"A1:{last_col}1")
    title_cell = ws["A1"]
    title_cell.value = "N梱包　教育完了リスト"
    title_cell.font = Font(name="游ゴシック", size=16, bold=True, color=COLOR_HEADER_FONT)
    title_cell.fill = PatternFill(start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # ── 凡例行 (row 2) ─────────────────────────────
    legend_col = get_column_letter(2 + len(MEMBERS))
    ws.merge_cells(f"A2:{legend_col}2")
    legend_cell = ws["A2"]
    legend_cell.value = "凡例：○ = 完了　　× = 未完了　　空白 = 未実施"
    legend_cell.font = Font(name="游ゴシック", size=9, color="595959")
    legend_cell.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 18

    # ── ヘッダー行 (row 3) ────────────────────────
    header_font   = Font(name="游ゴシック", size=11, bold=True, color=COLOR_HEADER_FONT)
    header_fill   = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
    header_align  = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 28

    headers = ["No.", "教育項目"] + MEMBERS
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = header_border()

    # ── データ行 (row 4〜) ────────────────────────
    item_font    = Font(name="游ゴシック", size=10)
    item_font_b  = Font(name="游ゴシック", size=10, bold=True)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align   = Alignment(horizontal="left",   vertical="center", indent=1)

    first_data_row = 4
    last_data_row  = first_data_row + len(ITEMS) - 1

    for row_offset, item in enumerate(ITEMS):
        row_num = first_data_row + row_offset
        ws.row_dimensions[row_num].height = 22

        # 行背景色（交互）
        row_bg = COLOR_ROW_ODD if row_offset % 2 == 0 else COLOR_ROW_EVEN
        row_fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")

        # No. 列
        no_cell = ws.cell(row=row_num, column=1, value=row_offset + 1)
        no_cell.font = item_font
        no_cell.fill = row_fill
        no_cell.alignment = center_align
        no_cell.border = thin_border()

        # 項目名列
        item_cell = ws.cell(row=row_num, column=2, value=item)
        item_cell.font = item_font_b
        item_cell.fill = row_fill
        item_cell.alignment = left_align
        item_cell.border = thin_border()

        # メンバー列（空白でデータ入力待ち）
        for col_offset in range(len(MEMBERS)):
            cell = ws.cell(row=row_num, column=3 + col_offset, value="")
            cell.fill = row_fill
            cell.alignment = center_align
            cell.border = thin_border()
            cell.font = Font(name="游ゴシック", size=12, bold=True)

    # ── 合計行 ──────────────────────────────────
    total_row = last_data_row + 1
    ws.row_dimensions[total_row].height = 24

    total_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    total_font = Font(name="游ゴシック", size=10, bold=True, color="276221")

    label_cell = ws.cell(row=total_row, column=1, value="完了数")
    label_cell.font = total_font
    label_cell.fill = total_fill
    label_cell.alignment = center_align
    label_cell.border = thin_border()

    ws.merge_cells(f"A{total_row}:B{total_row}")
    ws["A" + str(total_row)].value = "完了数（○の数）"
    ws["A" + str(total_row)].font = total_font
    ws["A" + str(total_row)].fill = total_fill
    ws["A" + str(total_row)].alignment = left_align

    for col_offset, _ in enumerate(MEMBERS):
        col = 3 + col_offset
        col_letter = get_column_letter(col)
        formula = f'=COUNTIF({col_letter}{first_data_row}:{col_letter}{last_data_row},"○")'
        cell = ws.cell(row=total_row, column=col, value=formula)
        cell.font = total_font
        cell.fill = total_fill
        cell.alignment = center_align
        cell.border = thin_border()

    # ── 条件付き書式（全データ範囲）────────────────
    data_range = f"C{first_data_row}:{get_column_letter(2 + len(MEMBERS))}{last_data_row}"
    apply_conditional_formatting(ws, data_range)

    # ── ウィンドウ枠の固定（ヘッダー + 項目列）────────
    ws.freeze_panes = "C4"

    return ws


def main():
    wb = Workbook()
    create_sheet(wb)
    wb.save(OUTPUT_FILE)
    print(f"作成完了: {OUTPUT_FILE}")
    print(f"  メンバー: {', '.join(MEMBERS)}")
    print(f"  教育項目: {len(ITEMS)} 項目")
    print("")
    print("セルへの入力方法:")
    print("  ○ → 緑色（教育完了）")
    print("  × → 赤色（教育未完了）")
    print("  空白 → 未実施")


if __name__ == "__main__":
    main()
