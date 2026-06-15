"""
LPN分割 誤り対応フロー（例外処理手順）生成スクリプト

既存の標準手順書フォーマット（手順No. / 手順 / 時間(秒) / 急所 / 写真）を
崩さずに、5パターンの場合分けを1枚のシートで表現する。

使用方法:
    python create_lpn_error_flow.py
    → LPN分割_誤り対応フロー.xlsx が生成されます
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUTPUT_FILE = "LPN分割_誤り対応フロー.xlsx"

# ── 色定義 ──────────────────────────────────────
COLOR_HEADER_BG   = "BFBFBF"   # グレー（見出し帯）
COLOR_DECISION    = "FFF2CC"   # 黄色（判定ステップ）
COLOR_CASE_12     = "DDEBF7"   # 青系（①②共通: OLPN統合）
COLOR_CASE_GREEN  = "E2EFDA"   # 緑系（ワンレック判定の先）
COLOR_CASE_3      = "E4DFEC"   # 紫系（③）
COLOR_CASE_4      = "FCE4D6"   # 橙系（④）
COLOR_CASE_5      = "D9D9D9"   # 灰系（⑤）

BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def cell(ws, row, col, value, *, fill=None, bold=False, size=11,
         wrap=True, align="left", valign="center", merge=None, font_color=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="游ゴシック", size=size, bold=bold, color=font_color)
    c.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
    c.border = BORDER
    if fill:
        c.fill = PatternFill(start_color=fill, end_color=fill, fill_type="solid")
    if merge:
        ws.merge_cells(start_row=row, start_column=col,
                        end_row=merge[0], end_column=merge[1])
        # 結合範囲の全セルに罫線を適用
        for r in range(row, merge[0] + 1):
            for cc in range(col, merge[1] + 1):
                ws.cell(row=r, column=cc).border = BORDER
                if fill:
                    ws.cell(row=r, column=cc).fill = PatternFill(
                        start_color=fill, end_color=fill, fill_type="solid")
    return c


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "LPN分割誤り対応"

    # ── 列幅 ──────────────────────────────────
    widths = {"A": 8, "B": 50, "C": 8, "D": 22, "E": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

    row = 1

    # ── タイトル ───────────────────────────────
    cell(ws, row, 1, "作業項目: 重長品検数作業 ／ LPN分割 誤り対応フロー（例外処理）",
         fill=COLOR_HEADER_BG, bold=True, size=13, align="center",
         merge=(row, 5))
    row += 1

    # ── ケース選択表（トリアージ表） ─────────────
    cell(ws, row, 1, "発生した誤り（症状）", fill=COLOR_HEADER_BG, bold=True,
         align="center", merge=(row, 2))
    cell(ws, row, 4, "対応ケース", fill=COLOR_HEADER_BG, bold=True, align="center")
    cell(ws, row, 5, "読み始める手順No.", fill=COLOR_HEADER_BG, bold=True, align="center")
    row += 1

    triage_rows = [
        ("① 即出荷後にLPN分割を忘れて即出荷してしまった", "①", "手順1から"),
        ("② LPN分割で数量を誤ってしまった", "②", "手順2から"),
        ("③ LPN分割を過剰に行ってしまった", "③", "手順6から"),
        ("④ 複数部材のLPN分割を途中で中断した", "④", "手順7から"),
        ("⑤ プリンタを設定しないままLPN分割した", "⑤", "手順8から"),
    ]
    for symptom, case_no, start_step in triage_rows:
        cell(ws, row, 1, symptom, merge=(row, 2))
        cell(ws, row, 4, case_no, align="center")
        cell(ws, row, 5, start_step, align="center")
        row += 1

    row += 1  # 空行

    # ── メイン手順テーブル ヘッダー ─────────────
    headers = ["手順No.", "手順", "時間(秒)", "急所", "写真"]
    for col_idx, h in enumerate(headers, start=1):
        cell(ws, row, col_idx, h, fill=COLOR_HEADER_BG, bold=True, align="center")
    row += 1

    # ── ステップ定義 ─────────────────────────────
    # (手順No, 手順テキスト, 急所テキスト, 写真テキスト, 色, 推奨行高)
    steps = [
        (
            "1",
            "【判定】即出荷後にLPN分割を実施したか？\n\n"
            "　YES → 手順2へ（OLPN統合）\n"
            "　NO  → 手順3へ（ワンレック判定）",
            "Q:\n出荷実績画面で「LPN分割」履歴の有無を確認する",
            "（画面例：出荷実績画面の分割履歴欄）",
            COLOR_DECISION,
            120,
        ),
        (
            "2",
            "【①YES／②共通】OLPN統合を実施する\n\n"
            "対象システム画面で対象LPNを選択し、統合処理を行う。",
            "S:\n統合対象のLPNを誤らないよう、現品票で再確認する",
            "（画面例：OLPN統合画面）",
            COLOR_CASE_12,
            110,
        ),
        (
            "3",
            "【判定】対象はワンレックか？\n\n"
            "　YES → 手順4へ（国内梱包：ワンレック）\n"
            "　NO  → 手順5へ（国内梱包：複数レック）",
            "Q:\n梱包指示書のレック数欄を確認する",
            "（画面例：梱包指示書）",
            COLOR_DECISION,
            110,
        ),
        (
            "4",
            "国内梱包（ワンレック）を実施する",
            "Q:\nレック内の品番・数量が指示と一致しているか確認する",
            "（画面例：ワンレック梱包時の画面）",
            COLOR_CASE_GREEN,
            80,
        ),
        (
            "5",
            "国内梱包（複数レック）を実施する",
            "Q:\n各レックの品番・数量が指示と一致しているか確認する",
            "（画面例：複数レック梱包時の画面）",
            COLOR_CASE_GREEN,
            80,
        ),
        (
            "6",
            "【③】分割ラベルを使用する\n\n"
            "LPN分割を過剰に行った場合、余分に発行されたラベルは"
            "「分割ラベル」として保管し、該当LPNに使用する。",
            "S:\n誤って別LPNに分割ラベルを貼付しないよう、ラベルの"
            "LPN番号を必ず確認する",
            "（画面例：分割ラベル保管箱）",
            COLOR_CASE_3,
            100,
        ),
        (
            "7",
            "【④】親部材集約を実施する\n\n"
            "複数部材のLPN分割作業を中断した場合、未処理の部材を"
            "親部材へ集約する。",
            "Q:\n集約後、親部材の数量が分割前の総数と一致しているか"
            "確認する",
            "（画面例：親部材集約画面）",
            COLOR_CASE_4,
            100,
        ),
        (
            "8",
            "【⑤】MAラベルを再印刷する\n\n"
            "プリンタを設定しないままLPN分割を行った場合、正しい"
            "プリンタを設定し、MAラベルを再印刷する。",
            "S:\n再印刷前に旧ラベルを破棄し、二重貼付を防止する",
            "（画面例：プリンタ設定画面／MAラベル印刷画面）",
            COLOR_CASE_5,
            100,
        ),
    ]

    for step_no, proc_text, key_text, photo_text, color, height in steps:
        cell(ws, row, 1, step_no, fill=color, bold=True, size=16, align="center")
        cell(ws, row, 2, proc_text, fill=color)
        cell(ws, row, 3, "", fill=color, align="center")
        cell(ws, row, 4, key_text, fill=color)
        cell(ws, row, 5, photo_text, fill=color, align="center", font_color="808080")
        ws.row_dimensions[row].height = height
        row += 1

    # ── 凡例 ─────────────────────────────────────
    row += 1
    cell(ws, row, 1, "凡例", fill=COLOR_HEADER_BG, bold=True, align="center")
    cell(ws, row, 2, "黄色＝判定ステップ（YES/NOで手順Noへ分岐） / "
                     "色帯＝対応ケース①〜⑤に対応", merge=(row, 5))
    row += 1

    ws.freeze_panes = "A9"

    wb.save(OUTPUT_FILE)
    print(f"作成完了: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
