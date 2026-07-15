"""
【即出荷】(2) 肉付け版 + 準備シート 生成スクリプト

方針:
  ・準備シート: タブレット・資材確認など事前準備のみ
  ・即出荷(2): 作業の「肝」(Q/S)に絞り、20手順→13手順に整理
  ・棚付け・oLPN分割への分岐を書き方1形式で記載
  ・書き方1ルール準拠: HGP創英角ゴシックUB / Arial Black / 白黒
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT_FILE = "即出荷2_肉付け版.xlsx"

# ── フォント（書き方1ルール）
FONT_NO   = Font(name="Arial Black",        size=24, bold=True)
FONT_PROC = Font(name="HGP創英角ｺﾞｼｯｸUB",  size=20)
FONT_SUPP = Font(name="HGP創英角ｺﾞｼｯｸUB",  size=14)
FONT_KEY  = Font(name="HGP創英角ｺﾞｼｯｸUB",  size=13)
FONT_HEAD = Font(name="HGP創英角ｺﾞｼｯｸUB",  size=11, bold=True)
FONT_BODY = Font(name="HGP創英角ｺﾞｼｯｸUB",  size=10)
FONT_BRANCH = Font(name="HGP創英角ｺﾞｼｯｸUB", size=13)

# ── 罫線
T = Side(style="thin")
H = Side(style="hair")
OUTER = Border(left=T, right=T, top=T, bottom=T)
INNER = Border(left=H, right=H, top=H, bottom=H)

# ── 配置
AL_R_TOP = Alignment(horizontal="right",  vertical="top",    wrap_text=True)
AL_L_TOP = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
AL_C_MID = Alignment(horizontal="center", vertical="center", wrap_text=True)
AL_L_MID = Alignment(horizontal="left",   vertical="center", wrap_text=True)

ROW_H = 22  # 行高(pt)


def sc(ws, coord, val=None, font=None, align=None, border=None, fill=None):
    c = ws[coord]
    if val   is not None: c.value = val
    if font  is not None: c.font  = font
    if align is not None: c.alignment = align
    if border is not None: c.border = border
    if fill  is not None: c.fill  = fill
    return c


def mg(ws, r1, c1, r2, c2, val=None, font=None, align=None, bdr=OUTER):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    coord = f"{get_column_letter(c1)}{r1}"
    sc(ws, coord, val, font, align, bdr)
    for r in range(r1, r2+1):
        for c in range(c1, c2+1):
            cell = ws.cell(row=r, column=c)
            if cell.border == Border():
                cell.border = Border(left=H, right=H, top=H, bottom=H)


def set_col_widths(ws, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def set_row_heights(ws, start, count, h=ROW_H):
    for i in range(start, start + count):
        ws.row_dimensions[i].height = h


# ── 1ステップ（12行）を描画
def build_step(ws, start, no, proc, supp="",
               s_text="", s_reason="", q_text="", q_reason=""):
    r, r6 = start, start + 6
    set_row_heights(ws, start, 12)

    mg(ws, r,  2, r+5,  2,  no,     FONT_NO,   AL_R_TOP)
    mg(ws, r,  3, r+5,  8,  proc,   FONT_PROC, AL_L_TOP)
    mg(ws, r6, 2, r6+5, 2,  "補足:", FONT_SUPP, AL_R_TOP)
    mg(ws, r6, 3, r6+5, 8,  supp,   FONT_SUPP, AL_L_TOP)
    mg(ws, r,  9, r+11, 9,  "",     None, None, INNER)

    mg(ws, r,   10, r+2,  10, "S:",    FONT_KEY, AL_R_TOP, INNER)
    mg(ws, r,   11, r+2,  13, s_text,  FONT_KEY, AL_L_TOP, INNER)
    mg(ws, r+3, 10, r+5,  10, "理由:", FONT_KEY, AL_R_TOP, INNER)
    mg(ws, r+3, 11, r+5,  13, s_reason, FONT_KEY, AL_L_TOP, INNER)
    mg(ws, r6,   10, r6+2, 10, "Q:",   FONT_KEY, AL_R_TOP, INNER)
    mg(ws, r6,   11, r6+2, 13, q_text,  FONT_KEY, AL_L_TOP, INNER)
    mg(ws, r6+3, 10, r6+5, 10, "理由:", FONT_KEY, AL_R_TOP, INNER)
    mg(ws, r6+3, 11, r6+5, 13, q_reason, FONT_KEY, AL_L_TOP, INNER)
    mg(ws, r,  14, r+11, 20, "", None, AL_C_MID)


# ══════════════════════════════════════════════
#  シート1: 準備シート
# ══════════════════════════════════════════════
def build_prep_sheet(wb):
    ws = wb.active
    ws.title = "【即出荷】準備"

    set_col_widths(ws, {"A": 4, "B": 6, "C": 40, "D": 30, "E": 20})

    FILL_H = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")
    FILL_ODD  = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    FILL_EVEN = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    # タイトル
    ws.row_dimensions[1].height = 30
    ws.merge_cells("A1:E1")
    sc(ws, "A1", "【即出荷】 事前準備チェックリスト",
       Font(name="HGP創英角ｺﾞｼｯｸUB", size=14, bold=True),
       AL_C_MID, OUTER, FILL_H)

    ws.row_dimensions[2].height = 20
    ws.merge_cells("A2:E2")
    sc(ws, "A2",
       "※この準備が完了したら【即出荷】(2) シートへ。棚付けをする場合は【棚付け】シートへ。",
       FONT_BODY, AL_L_MID)

    # ヘッダー
    ws.row_dimensions[3].height = 20
    for col, txt in zip(["A","B","C","D","E"],
                        ["", "No.", "確認項目", "確認ポイント", "備考"]):
        sc(ws, f"{col}3", txt,
           Font(name="HGP創英角ｺﾞｼｯｸUB", size=10, bold=True),
           AL_C_MID, OUTER, FILL_H)

    items = [
        ("1", "タブレットD（通過検品_即出荷）の電源ON",
         "ロック解除後、WMSアプリが起動しているか確認", "スキャナ接続も確認"),
        ("2", "スキャナ（タブレットD接続）の動作確認",
         "バーコードを読み取れるかテストスキャン", ""),
        ("3", "「納品書在中」スタンプとスタンプ台の準備",
         "インク残量確認", ""),
        ("4", "梱包資材の在庫確認",
         "出荷箱（各サイズ）・クラフトテープ・OPPテープ・緩衝材",
         "不足時はリーダーへ連絡"),
        ("5", "工程内不良カードの場所確認",
         "所定の棚にあるか確認", "【oLPN分割】シート参照"),
        ("6", "作業台（HP側 or 転送側）の決定",
         "担当割振りに従い使用エリアを確認", "転送側はドーリーも準備"),
    ]

    for i, (no, item, point, note) in enumerate(items):
        r = 4 + i
        ws.row_dimensions[r].height = 28
        fill = FILL_ODD if i % 2 == 0 else FILL_EVEN
        sc(ws, f"A{r}", "□",
           Font(name="游ゴシック", size=12), AL_C_MID, OUTER, fill)
        sc(ws, f"B{r}", no,
           FONT_BODY, AL_C_MID, OUTER, fill)
        sc(ws, f"C{r}", item,
           Font(name="HGP創英角ｺﾞｼｯｸUB", size=10), AL_L_MID, OUTER, fill)
        sc(ws, f"D{r}", point,
           FONT_BODY, AL_L_MID, OUTER, fill)
        sc(ws, f"E{r}", note,
           FONT_BODY, AL_L_MID, OUTER, fill)


# ══════════════════════════════════════════════
#  シート2: 【即出荷】(2) 本文
# ══════════════════════════════════════════════
def build_main_sheet(wb):
    ws = wb.create_sheet("【即出荷】(2)")

    set_col_widths(ws, {
        "A": 10.33, "B": 10.16, "C": 8.5,  "D": 8.5,  "E": 8.5,
        "F": 8.5,   "G": 8.5,   "H": 8.5,  "I": 10.33, "J": 6.33,
        "K": 8.33,  "L": 8.5,   "M": 8.5,  "N": 9.33,
    })

    FILL_HEAD = PatternFill(start_color="BFBFBF", end_color="BFBFBF", fill_type="solid")

    # ── ヘッダー行（元フォーマット準拠）
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 20
    ws.row_dimensions[6].height = 22

    mg(ws, 1, 2, 1, 7,  "標準作業手順書", FONT_HEAD, AL_C_MID, OUTER)
    mg(ws, 1, 8, 1, 13, "文書番号: 02-中RC-AP-02-0019", FONT_BODY, AL_L_MID)

    mg(ws, 2, 2, 2, 7,  "作業項目: 【即出荷】", FONT_HEAD, AL_L_MID)

    mg(ws, 3, 2, 5, 8,
       "【場合分け】\n"
       "・棚付けをする場合\n"
       "　　→ 作業項目【棚付け】→ 手順【1】へ\n"
       "・数量不一致が発生した場合\n"
       "　　→ 作業項目【oLPN分割】→ 工程内不良カード記載へ",
       FONT_BRANCH, AL_L_TOP, OUTER)

    # 列ヘッダー
    for col, txt in [(2,"手順No."),(3,"手順"),(9,"時間(秒)"),(10,"急所")]:
        sc(ws, f"{get_column_letter(col)}6", txt, FONT_HEAD, AL_C_MID, OUTER, FILL_HEAD)
    mg(ws, 6, 3, 6, 8,  "手順", FONT_HEAD, AL_C_MID, OUTER)
    mg(ws, 6, 10, 6, 13, "急所", FONT_HEAD, AL_C_MID, OUTER)
    mg(ws, 6, 14, 6, 20, "写真", FONT_HEAD, AL_C_MID, OUTER)

    row = 7

    # ── 手順1: ラベル確認（分岐起点）
    build_step(
        ws, row, "1",
        "ラベルを確認する",
        "部材の場合、1つのオリコンに複数ラベルが貼ってある。\nすべてのラベルを確認してから作業開始すること。\n\n"
        "・棚付けをする場合　→ 【棚付け】手順【1】へ\n"
        "・即出荷の場合　　　→ 手順【2】へ",
        q_text="ラベル枚数・発送先・品番・数量の確認漏れ厳禁",
        q_reason="ラベル見落としは誤出荷の直接原因になるため",
    )
    row += 12

    # ── 手順2: オリコン移動・開閉
    build_step(
        ws, row, "2",
        "オリコンを作業台に移動し、開閉して商品を取り出す",
        "HP側：オリコンを作業台に引き込む\n"
        "転送側：ドーリーを持ってきて作業台にオリコンを置く\n"
        "開閉手順は付表①シート参照。必ず両手・置いた状態で開閉する。\n"
        "封が閉まっている箱・袋は開梱せずそのまま使用する。",
        s_text="オリコン開閉時、空中作業・片手作業は禁止",
        s_reason="指を挟む危険があるため（過去に挟み事故あり）",
        q_text="封が閉まっている箱・シーラ袋はそのまま使用（開梱して取り出さない）",
        q_reason="開梱時の商品落下→数量違い発生のため",
    )
    row += 12

    # ── 手順3: 商品を梱包台へ（西側・長物注意）
    build_step(
        ws, row, "3",
        "商品を1点ずつ梱包台に置く",
        "西側・転送作業場での作業時、商品を1点ずつ梱包台に置く。\n"
        "西側での長物品は商品を押さえながら取ること。\n\n"
        "作業台の上は必ず1オーダー分のみ（同じお客様名でも送り先が違う場合は混在禁止）。",
        s_text="西側：倒れやすい長物品は商品を押さえて取る",
        s_reason="過去に長尺品がバランスを崩し目に入る事故が発生したため",
        q_text="作業台は1オーダー分のみ。送り先の異なる商品の混在禁止",
        q_reason="他のお客様の商品が混入するのを防ぐため",
    )
    row += 12

    # ── 手順4: 箱作成・テープ貼り
    build_step(
        ws, row, "4",
        "商品サイズに合った出荷箱を選定し、底面にクラフトテープを貼る",
        "目安: オリコン ≒ No.8箱（同サイズ）\n"
        "①十字張り → 片手で持てない重量品\n"
        "②H張り   → No.6以上の大箱（1辺の半分以上の長さを使用）",
        q_text="商品重量・サイズに応じたテープの貼り方を選定",
        q_reason="貼り方不足による底抜け→商品破損につながるため",
    )
    row += 12

    # ── 手順5: 緩衝材＋商品箱詰め＋カウント
    build_step(
        ws, row, "5",
        "底面に緩衝材を敷き、カウントしながら商品を箱に入れる",
        "特性フラグを確認すること。\n"
        "・パック品：2ピース以上のまとまり → バラさずカウント\n"
        "・バラ品：バラし漏れ注意、1ピースずつ確認\n"
        "・セット品：構成品が揃っているか確認\n"
        "メーカー箱の場合: 箱側面のメーカー入数を必ず確認し、赤丸をつける。",
        q_text="底面・側面が隠れるよう緩衝材を敷く。入数情報を確認し赤丸をつける",
        q_reason="緩衝材不足→破損 / 赤丸なし→数量確認抜け漏れ防止のため",
    )
    row += 12

    # ── 手順6: 数量確認・赤丸
    build_step(
        ws, row, "6",
        "カウント数がピックラベルの数量と一致しているか確認し、赤丸をつける",
        "一致しない場合 → 工程内不良カードを記入し、管理社員へ渡す\n"
        "　（→ 作業項目【oLPN分割】→ 工程内不良カード記載へ）\n\n"
        "工程内不良カード記入内容:\n"
        "①日付 ②不良内容 ③④品番・ラベル情報 ⑤自身の氏名",
        q_text="部材品は並べてカウント。点数と番号を両方チェック",
        q_reason="1個単位の数量違いが最多ミスのため、並べることで視認精度を上げる",
    )
    row += 12

    # ── 手順7: スキャン
    build_step(
        ws, row, "7",
        "タブレットD（通過検品_即出荷）のスキャナで\nG#（OLPN）ラベルをスキャンする",
        "20kg以上の商品・S18箱にしか入らない商品は\n"
        "配送方法をQ/T（キャリー/ヤマト）→ S（佐川）へ切り替えること。",
    )
    row += 12

    # ── 手順8: 発行確認（シューター番号）
    build_step(
        ws, row, "8",
        "発行された納品書・送り状・ランクラベルを受け取り確認する",
        "①納品書  → 残1梱包のみ出力（複数枚出ることあり）\n"
        "②送り状  → シューター番号を納品書と照合\n"
        "③ランクラベル → ランクオーダーのみ出力",
        q_text="納品書と送り状のシューター番号が一致しているか確認",
        q_reason="印刷機トラブル等で直前の別注文のものと混在しないため（誤配送防止）",
    )
    row += 12

    # ── 手順9: 送り状貼付
    build_step(
        ws, row, "9",
        "送り状を箱の長手面・右上に貼り付ける",
        "箱の長い面（長手面）の右上に貼る。\n"
        "広告印刷がある場合は、なるべく見える位置を選ぶ。",
        q_text="送り状の貼り忘れ厳禁",
        q_reason="送り状がないと荷合せ不能 / お客様が届いた商品の内容不明になるため",
    )
    row += 12

    # ── 手順10: 緩衝材詰め・封緘（最重要Q）
    build_step(
        ws, row, "10",
        "上部まで緩衝材を詰めて封緘（クラフトテープ）を行う",
        "袋の場合：開封口を1回折り込みOPPテープで止め、その後箱に入れる\n"
        "箱の場合：天面をクラフトテープで封緘\n"
        "封緘後に必ず箱を振って商品が動かないか確認する。",
        q_text="封緘後に箱を振って、中で商品が動かないか確認する。破損箱は交換",
        q_reason="輸送中の振動・衝撃による商品・箱の破損防止のため",
    )
    row += 12

    # ── 手順11: 納品書在中スタンプ・ランクラベル
    build_step(
        ws, row, "11",
        "「納品書在中」印を捺印し、ランクラベルを商品ラベルの左隣に貼る",
        "「納品書在中」スタンプ：送り状から見て左上、広告と被らない位置に捺印\n"
        "ランクラベル：商品ラベルの左隣に貼付\n\n"
        "・特定伝(O1)・指定伝(O2-O6)の場合\n"
        "　→ 封緘を再度実施し、特定指定伝N品置場へ搬送（手順12へ）\n"
        "・上記以外の場合 → 手順13へ",
    )
    row += 12

    # ── 手順12: 特定指定伝
    build_step(
        ws, row, "12",
        "【特定伝・指定伝のみ】封緘を実施し、所定置場へ搬送する",
        "手順10と同じ封緘作業を実施する。\n"
        "搬送先: 特定指定伝N品置場",
    )
    row += 12

    # ── 手順13: 搬送
    build_step(
        ws, row, "13",
        "出荷箱を所定の置場へ搬送する",
        "前当(前日注文当日出荷)の場合 → 3F 入出荷エリアへ\n"
        "当1・当2・当3の場合 → 4F 指定伝梱包台へ",
    )


def main():
    wb = Workbook()
    build_prep_sheet(wb)
    build_main_sheet(wb)
    wb.save(OUTPUT_FILE)
    print(f"作成完了: {OUTPUT_FILE}")
    print("  シート1: 【即出荷】準備（事前準備チェックリスト）")
    print("  シート2: 【即出荷】(2)（肝に絞った13手順版）")


if __name__ == "__main__":
    main()
