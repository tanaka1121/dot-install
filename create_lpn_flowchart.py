"""
LPN分割 誤り対応フロー図 - Excel ネイティブ図形版
  ・開始  : 角丸四角形 (roundRect, 高丸め) ← 楕円□
  ・分岐  : ひし形 (diamond) ◇
  ・処理  : 四角形 (rect) □
  ・発生事象: 角丸四角形・点線枠
  ・終了ノードなし
  ・すべての図形が Excel 上で直接編集可能
"""

import zipfile, io, re
from openpyxl import Workbook

OUTPUT = "LPN分割_誤り対応フロー図.xlsx"
CM = 360000   # 1cm → EMU

# ── 色 ──────────────────────────────────────────────────────────────
START_FILL = "BDD7EE"   # 角丸四角（開始）薄青
DEC_FILL   = "FFE699"   # ひし形（分岐）
ACT_FILL   = "DEEAF1"   # 四角（処理）
CASE_FILL  = "FFFFFF"   # 発生事象（白・点線）
LINE_COL   = "44546A"   # 罫線・矢印
YES_COL    = "375623"   # YES ラベル
NO_COL     = "9C0006"   # NO ラベル


def e(v): return int(v * CM)          # cm → EMU
def ep(*v): return tuple(e(x) for x in v)  # tuple convert


# ──────────────────────────────────────────────────────────────────────
# XML パーツ生成
# ──────────────────────────────────────────────────────────────────────
_sid = [1]

def next_id():
    _sid[0] += 1
    return _sid[0]


def _anchor(x, y, w, h, inner_xml):
    """oneCellAnchor wrapper"""
    return f"""<xdr:oneCellAnchor>
  <xdr:from><xdr:col>0</xdr:col><xdr:colOff>{x}</xdr:colOff>
             <xdr:row>0</xdr:row><xdr:rowOff>{y}</xdr:rowOff></xdr:from>
  <xdr:ext cx="{w}" cy="{h}"/>
  {inner_xml}
  <xdr:clientData/>
</xdr:oneCellAnchor>"""


def _text_paras(lines, sz, bold=False, color="000000"):
    b = "<a:b/>" if bold else ""
    paras = []
    for ln in lines:
        paras.append(
            f'<a:p><a:pPr algn="ctr"/>'
            f'<a:r><a:rPr lang="ja-JP" sz="{sz}" dirty="0" b="{1 if bold else 0}">'
            f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill></a:rPr>'
            f'<a:t>{ln}</a:t></a:r></a:p>'
        )
    return "\n".join(paras)


def shape(x, y, w, h, text, prst, fill, adj=None,
          dashed=False, font_sz=1000, bold=False, line_w=25400):
    """図形 XML を返す (oneCellAnchor)"""
    sid = next_id()
    adj_xml = f'<a:avLst><a:gd name="adj" fmla="val {adj}"/></a:avLst>' if adj else "<a:avLst/>"
    dash_xml = '<a:prstDash val="dash"/>' if dashed else ""
    lines = text.split("\n")
    paras = _text_paras(lines, font_sz, bold)
    inner = f"""<xdr:sp macro="" textlink="">
  <xdr:nvSpPr>
    <xdr:cNvPr id="{sid}" name="Shape{sid}"/>
    <xdr:cNvSpPr><a:spLocks noGrp="1"/></xdr:cNvSpPr>
  </xdr:nvSpPr>
  <xdr:spPr>
    <a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>
    <a:prstGeom prst="{prst}">{adj_xml}</a:prstGeom>
    <a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>
    <a:ln w="{line_w}"><a:solidFill><a:srgbClr val="{LINE_COL}"/></a:solidFill>{dash_xml}</a:ln>
  </xdr:spPr>
  <xdr:txBody>
    <a:bodyPr wrap="square" anchor="ctr"><a:normAutofit/></a:bodyPr>
    <a:lstStyle/>
    {paras}
  </xdr:txBody>
</xdr:sp>"""
    return _anchor(x, y, w, h, inner)


def connector(x1, y1, x2, y2):
    """矢印付き直線コネクタ (straightConnector1)"""
    sid = next_id()
    dx, dy = x2 - x1, y2 - y1
    flip = ""
    # bounding box top-left + extents
    bx = min(x1, x2)
    by = min(y1, y2)
    cx_ = max(abs(dx), 9525)   # 最小 1px
    cy_ = max(abs(dy), 9525)
    if dx < 0:  flip += ' flipH="1"'
    if dy < 0:  flip += ' flipV="1"'
    inner = f"""<xdr:cxnSp macro="">
  <xdr:nvCxnSpPr>
    <xdr:cNvPr id="{sid}" name="Conn{sid}"/>
    <xdr:cNvCxnSpPr/>
  </xdr:nvCxnSpPr>
  <xdr:spPr>
    <a:xfrm{flip}><a:off x="0" y="0"/><a:ext cx="{cx_}" cy="{cy_}"/></a:xfrm>
    <a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>
    <a:ln w="25400">
      <a:solidFill><a:srgbClr val="{LINE_COL}"/></a:solidFill>
      <a:tailEnd type="arrow" w="med" len="med"/>
    </a:ln>
  </xdr:spPr>
</xdr:cxnSp>"""
    return _anchor(bx, by, cx_, cy_, inner)


def label(x, y, w, h, text, color=YES_COL):
    """YES/NO ラベル用テキストボックス"""
    sid = next_id()
    paras = _text_paras([text], sz=900, bold=True, color=color)
    inner = f"""<xdr:sp macro="" textlink="">
  <xdr:nvSpPr>
    <xdr:cNvPr id="{sid}" name="Lbl{sid}"/>
    <xdr:cNvSpPr txBox="1"><a:spLocks noGrp="1"/></xdr:cNvSpPr>
  </xdr:nvSpPr>
  <xdr:spPr>
    <a:xfrm><a:off x="0" y="0"/><a:ext cx="{w}" cy="{h}"/></a:xfrm>
    <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
    <a:noFill/>
    <a:ln><a:noFill/></a:ln>
  </xdr:spPr>
  <xdr:txBody>
    <a:bodyPr wrap="square" anchor="ctr"><a:normAutofit/></a:bodyPr>
    <a:lstStyle/>
    {paras}
  </xdr:txBody>
</xdr:sp>"""
    return _anchor(x, y, w, h, inner)


def separator(y):
    """シナリオ間の区切り横線"""
    sid = next_id()
    inner = f"""<xdr:cxnSp macro="">
  <xdr:nvCxnSpPr>
    <xdr:cNvPr id="{sid}" name="Sep{sid}"/>
    <xdr:cNvCxnSpPr/>
  </xdr:nvCxnSpPr>
  <xdr:spPr>
    <a:xfrm><a:off x="0" y="0"/><a:ext cx="{e(13)}" cy="{9525}"/></a:xfrm>
    <a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>
    <a:ln w="9525" cmpd="sng">
      <a:solidFill><a:srgbClr val="BFBFBF"/></a:solidFill>
      <a:prstDash val="dash"/>
    </a:ln>
  </xdr:spPr>
</xdr:cxnSp>"""
    return _anchor(e(0.3), y, e(13), 9525, inner)


# ──────────────────────────────────────────────────────────────────────
# フローチャート定義
# ──────────────────────────────────────────────────────────────────────
def build_drawing():
    parts = []

    # ── ショートカット ──────────────────────────────────────────────
    def S(x, y, w, h, txt, **kw):   return shape(e(x),e(y),e(w),e(h),txt,**kw)
    def C(x1,y1,x2,y2):             return connector(e(x1),e(y1),e(x2),e(y2))
    def L(x,y,w,h,txt,col=YES_COL): return label(e(x),e(y),e(w),e(h),txt,col)
    def SEP(y):                      return separator(e(y))

    # ════════════════════════════════════════════════════════════════
    # ① LPN分割を忘れて即出荷してしまった
    # ════════════════════════════════════════════════════════════════
    # 開始①（角丸四角・楕円□）
    parts.append(S(0.3, 0.2,  3.0, 0.65,
                   "開始①", prst="roundRect", fill=START_FILL, adj=50000, font_sz=1000, bold=True))
    # ↓
    parts.append(C(1.8, 0.85, 1.8, 1.2))
    # 発生事象ボックス（点線）
    parts.append(S(0.3, 1.2,  3.0, 1.1,
                   "①LPN分割を忘れて\n即出荷してしまった",
                   prst="roundRect", fill=CASE_FILL, dashed=True, font_sz=950))
    # → Decision1
    parts.append(C(3.3, 1.75, 4.3, 1.75))

    # Decision1 ◇（即出荷後にLPN分割したか？）
    parts.append(S(4.3, 1.1,  3.6, 1.3,
                   "即出荷後に\nLPN分割したか？",
                   prst="diamond", fill=DEC_FILL, font_sz=950))
    # YES→ 右
    parts.append(C(7.9, 1.75, 9.0, 1.75))
    parts.append(L(7.95, 1.45, 0.8, 0.4, "YES"))
    # OLPNを統合する □
    parts.append(S(9.0, 1.3,  3.5, 0.9,
                   "OLPNを統合する",
                   prst="rect", fill=ACT_FILL, font_sz=1000))

    # NO↓
    parts.append(C(6.1, 2.4, 6.1, 3.1))
    parts.append(L(6.2, 2.6, 0.7, 0.35, "NO", col=NO_COL))

    # Decision2 ◇（ワンレックか？）
    parts.append(S(4.3, 3.1,  3.6, 1.3,
                   "対象は\nワンレックか？",
                   prst="diamond", fill=DEC_FILL, font_sz=950))
    # YES→ 右
    parts.append(C(7.9, 3.75, 9.0, 3.75))
    parts.append(L(7.95, 3.45, 1.0, 0.4, "ワンレック"))
    # 国内梱包（ワンレック）
    parts.append(S(9.0, 3.3,  3.5, 0.9,
                   "国内梱包\n（ワンレック）",
                   prst="rect", fill=ACT_FILL, font_sz=1000))

    # NO↓
    parts.append(C(6.1, 4.4, 6.1, 5.1))
    parts.append(L(6.2, 4.55, 1.1, 0.35, "複数レック", col=NO_COL))
    # 国内梱包（複数レック）
    parts.append(S(4.3, 5.1,  3.6, 1.1,
                   "国内梱包（複数レック）\n＋イレギュラー置場へ",
                   prst="rect", fill=ACT_FILL, font_sz=950))

    # ════════════════════════════════════════════════════════════════
    # ② LPN分割で数量を誤った
    # ════════════════════════════════════════════════════════════════
    parts.append(SEP(7.0))
    parts.append(S(0.3, 7.2,  3.0, 0.65,
                   "開始②", prst="roundRect", fill=START_FILL, adj=50000, font_sz=1000, bold=True))
    parts.append(C(1.8, 7.85, 1.8, 8.2))
    parts.append(S(0.3, 8.2,  3.0, 0.9,
                   "②LPN分割で\n数量を誤った",
                   prst="roundRect", fill=CASE_FILL, dashed=True, font_sz=950))
    parts.append(C(3.3, 8.65, 4.5, 8.65))
    parts.append(S(4.5, 8.2,  3.5, 0.9,
                   "OLPNを統合する",
                   prst="rect", fill=ACT_FILL, font_sz=1000))

    # ════════════════════════════════════════════════════════════════
    # ③ LPN分割を過剰に行った
    # ════════════════════════════════════════════════════════════════
    parts.append(SEP(10.0))
    parts.append(S(0.3, 10.2, 3.0, 0.65,
                   "開始③", prst="roundRect", fill=START_FILL, adj=50000, font_sz=1000, bold=True))
    parts.append(C(1.8, 10.85, 1.8, 11.2))
    parts.append(S(0.3, 11.2, 3.0, 0.9,
                   "③LPN分割を\n過剰に行った",
                   prst="roundRect", fill=CASE_FILL, dashed=True, font_sz=950))
    parts.append(C(3.3, 11.65, 4.5, 11.65))
    parts.append(S(4.5, 11.2, 3.5, 0.9,
                   "分割ラベルを使用する",
                   prst="rect", fill=ACT_FILL, font_sz=1000))

    # ════════════════════════════════════════════════════════════════
    # ④ 複数部材のLPN分割を途中で中断した
    # ════════════════════════════════════════════════════════════════
    parts.append(SEP(13.0))
    parts.append(S(0.3, 13.2, 3.0, 0.65,
                   "開始④", prst="roundRect", fill=START_FILL, adj=50000, font_sz=1000, bold=True))
    parts.append(C(1.8, 13.85, 1.8, 14.2))
    parts.append(S(0.3, 14.2, 3.0, 1.0,
                   "④複数部材のLPN分割を\n途中で中断した",
                   prst="roundRect", fill=CASE_FILL, dashed=True, font_sz=950))
    parts.append(C(3.3, 14.7, 4.5, 14.7))
    parts.append(S(4.5, 14.2, 3.5, 0.9,
                   "親部材集約を実施する",
                   prst="rect", fill=ACT_FILL, font_sz=1000))

    # ════════════════════════════════════════════════════════════════
    # ⑤ プリンタを設定しないままLPN分割した
    # ════════════════════════════════════════════════════════════════
    parts.append(SEP(16.0))
    parts.append(S(0.3, 16.2, 3.0, 0.65,
                   "開始⑤", prst="roundRect", fill=START_FILL, adj=50000, font_sz=1000, bold=True))
    parts.append(C(1.8, 16.85, 1.8, 17.2))
    parts.append(S(0.3, 17.2, 3.0, 1.0,
                   "⑤プリンタを設定しない\nままLPN分割した",
                   prst="roundRect", fill=CASE_FILL, dashed=True, font_sz=950))
    parts.append(C(3.3, 17.7, 4.5, 17.7))
    parts.append(S(4.5, 17.2, 3.5, 0.9,
                   "MAラベルを再印刷する",
                   prst="rect", fill=ACT_FILL, font_sz=1000))

    NS = (
        'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    )
    return (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<xdr:wsDr {NS}>\n'
        + "\n".join(parts)
        + "\n</xdr:wsDr>"
    )


# ──────────────────────────────────────────────────────────────────────
# xlsx パッチ（openpyxl で保存後、ZIP に描画 XML を注入）
# ──────────────────────────────────────────────────────────────────────
SHEET_REL_ID = "rId10"
DRAWING_PATH = "xl/drawings/drawing1.xml"
DRAWING_PART_NAME = "/xl/drawings/drawing1.xml"
CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.drawing+xml'
REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"


def patch_xlsx(wb, drawing_xml: str, output_path: str):
    # 1. openpyxl で基本 xlsx をバッファに保存
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    # 2. ZIP を読み込んで修正
    out_buf = io.BytesIO()
    with zipfile.ZipFile(buf, "r") as zin, zipfile.ZipFile(out_buf, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for name in names:
            data = zin.read(name)

            # [Content_Types].xml に描画の ContentType を追加
            if name == "[Content_Types].xml":
                text = data.decode("utf-8")
                insert = f'<Override PartName="{DRAWING_PART_NAME}" ContentType="{CONTENT_TYPE}"/>'
                text = text.replace("</Types>", insert + "</Types>")
                data = text.encode("utf-8")

            # ワークシート XML に <drawing> 参照を追加
            elif name == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                drawing_ref = f'<drawing r:id="{SHEET_REL_ID}"/>'
                if drawing_ref not in text:
                    text = text.replace("</worksheet>", drawing_ref + "</worksheet>")
                data = text.encode("utf-8")

            # ワークシートのリレーション追加 or 作成
            elif name == "xl/worksheets/_rels/sheet1.xml.rels":
                text = data.decode("utf-8")
                rel = (f'<Relationship Id="{SHEET_REL_ID}" Type="{REL_TYPE}" '
                       f'Target="../drawings/drawing1.xml"/>')
                text = text.replace("</Relationships>", rel + "</Relationships>")
                data = text.encode("utf-8")

            zout.writestr(name, data)

        # ワークシートの _rels が存在しない場合は作成
        rels_path = "xl/worksheets/_rels/sheet1.xml.rels"
        if rels_path not in names:
            rel_content = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                f'<Relationship Id="{SHEET_REL_ID}" Type="{REL_TYPE}" '
                f'Target="../drawings/drawing1.xml"/>'
                '</Relationships>'
            )
            zout.writestr(rels_path, rel_content.encode("utf-8"))

        # 3. 描画 XML を追加
        zout.writestr(DRAWING_PATH, drawing_xml.encode("utf-8"))

    # 4. ファイルに書き出し
    out_buf.seek(0)
    with open(output_path, "wb") as f:
        f.write(out_buf.read())


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "LPN分割誤り対応フロー"
    ws["A1"] = "LPN分割 誤り対応フロー図"
    ws["A1"].font = __import__("openpyxl").styles.Font(
        name="HGP創英角ｺﾞｼｯｸUB", size=16, bold=True)
    ws.row_dimensions[1].height = 20

    drawing_xml = build_drawing()
    patch_xlsx(wb, drawing_xml, OUTPUT)
    print(f"作成完了: {OUTPUT}")
    print("図形はすべて Excel 上でクリックして直接編集できます。")


if __name__ == "__main__":
    main()
